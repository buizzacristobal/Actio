import json
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.api.v1.endpoints.extraction import get_llm_client
from app.core.config import Settings, get_settings
from app.main import create_app

PARTY_C = {"name": "Banco Ejemplo SpA", "rut": "76.086.428-5"}
PARTY_D = {"name": "Juan Pérez", "rut": "12.345.678-5"}
PAGARE = {"kind": "pagare", "amount": "1000000", "currency": "CLP",
          "issue_date": "2024-01-20", "maturity_date": "2024-03-20", "agreed_rate_annual": "24",
          "creditor": PARTY_C, "debtor": PARTY_D,
          "jurisdiction": {"court_name": "12° Juzgado Civil", "city": "Santiago"}}
ATT = {"name": "María Soto", "rut": "11.111.111-1", "address": "Huérfanos 1117, Santiago"}


class Fake:
    def __init__(self, reply): self.reply = reply
    def complete(self, system, user): return self.reply


@pytest.fixture
def app():
    return create_app()


@pytest.fixture
def client(app):
    return TestClient(app, raise_server_exceptions=False)


def test_health(client):
    j = client.get("/health").json()
    assert j["status"] == "ok" and j["version"]


def test_extract_ok_and_bad_output(app, client):
    app.dependency_overrides[get_llm_client] = lambda: Fake(json.dumps(PAGARE))
    r = client.post("/api/v1/extract", json={"text": "PAGARE ..."})
    assert r.status_code == 200 and r.json()["instrument"]["kind"] == "pagare"
    app.dependency_overrides[get_llm_client] = lambda: Fake("sin json")
    r = client.post("/api/v1/extract", json={"text": "SECRETO-DOC"})
    assert r.status_code == 422 and "SECRETO-DOC" not in r.text


def test_liquidate(client):
    body = {"principal": "1000000", "currency": "CLP", "issue_date": "2024-01-20",
            "maturity_date": "2024-03-20", "cutoff_date": "2024-04-20", "day_count": "ACT/365"}
    r = client.post("/api/v1/liquidate", json=body)
    assert r.status_code == 200
    j = r.json()
    assert j["category"] == "non_reajustable_clp_under_50_uf" and j["day_count"] == "ACT/365"
    assert int(j["total_clp"]) == 1_000_000 + int(j["interest_clp"]) and len(j["days"]) == 91
    assert sum(int(m["interest"]) for m in j["months"]) == int(j["interest_native"])
    assert any("NO verificadas" in w for w in j["warnings"])


def test_liquidate_errors_do_not_echo_input(client):
    bad = {"principal": "-5", "currency": "CLP", "issue_date": "2024-01-20",
           "maturity_date": "2024-03-20", "cutoff_date": "2024-04-20"}
    r = client.post("/api/v1/liquidate", json=bad)
    assert r.status_code == 422 and "-5" not in r.text
    out_of_range = {**bad, "principal": "1000", "issue_date": "2010-01-01",
                    "maturity_date": "2010-01-02", "cutoff_date": "2010-02-01"}
    assert client.post("/api/v1/liquidate", json=out_of_range).status_code == 422


def test_generate_lawsuit_docx(client):
    r = client.post("/api/v1/generate-lawsuit", json={
        "instrument": PAGARE, "attorney": ATT, "cutoff_date": "2024-04-20"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "demanda_ejecutiva.docx" in r.headers["content-disposition"]
    doc = Document(BytesIO(r.content))
    txt = "\n".join(p.text for p in doc.paragraphs)
    assert "BORRADOR" in txt and "S.J.L. EN LO CIVIL DE SANTIAGO" in txt and doc.tables


def test_generate_lawsuit_requires_jurisdiction_and_factura_due_date(client):
    no_court = {k: v for k, v in PAGARE.items() if k != "jurisdiction"}
    r = client.post("/api/v1/generate-lawsuit", json={
        "instrument": no_court, "attorney": ATT, "cutoff_date": "2024-04-20"})
    assert r.status_code == 422
    fac = {"kind": "factura", "folio": "1", "net_amount": "100", "vat_amount": "19",
           "total_amount": "119", "issue_date": "2024-01-20", "creditor": PARTY_C,
           "debtor": PARTY_D, "jurisdiction": PAGARE["jurisdiction"]}
    r = client.post("/api/v1/generate-lawsuit", json={
        "instrument": fac, "attorney": ATT, "cutoff_date": "2024-04-20"})
    assert r.status_code == 422 and "due_date" in r.text


def test_api_key_enforced_when_configured(app):
    app.dependency_overrides[get_settings] = lambda: Settings(api_key="k")
    c = TestClient(app)
    body = {"principal": "1000", "currency": "CLP", "issue_date": "2024-01-20",
            "maturity_date": "2024-01-21", "cutoff_date": "2024-01-22"}
    assert c.post("/api/v1/liquidate", json=body).status_code == 401
    assert c.post("/api/v1/liquidate", json=body, headers={"X-API-Key": "k"}).status_code == 200


def test_hardening_headers_rate_limit_and_body_cap(monkeypatch):
    monkeypatch.setenv("RATE_LIMIT_PER_MIN", "3")
    monkeypatch.setenv("MAX_BODY_BYTES", "500")
    c = TestClient(create_app())
    r = c.get("/health")
    assert r.headers["x-content-type-options"] == "nosniff" and r.headers["x-request-id"]
    body = {"principal": "1000", "currency": "CLP", "issue_date": "2024-01-20",
            "maturity_date": "2024-01-21", "cutoff_date": "2024-01-22"}
    codes = [c.post("/api/v1/liquidate", json=body).status_code for _ in range(4)]
    assert codes == [200, 200, 200, 429]
    big = TestClient(create_app()).post("/api/v1/extract", json={"text": "x" * 2000})
    assert big.status_code == 413


def test_preview_and_edited_generation(client):
    body = {"instrument": PAGARE, "attorney": ATT, "cutoff_date": "2024-04-20",
            "assets": ["Vehículo PPU ABCD12"]}
    r = client.post("/api/v1/preview-lawsuit", json=body)
    assert r.status_code == 200
    j = r.json()
    assert j["blocks"][0]["kind"] == "banner" and j["notes"]
    assert any("Vehículo PPU ABCD12" in b["text"] for b in j["blocks"])
    edited = [{**b, "text": b["text"] + " [rev]"} if b["kind"] == "paragraph" else b for b in j["blocks"]]
    r = client.post("/api/v1/generate-lawsuit", json={**body, "edited_blocks": edited})
    assert r.status_code == 200
    txt = "\n".join(p.text for p in Document(BytesIO(r.content)).paragraphs)
    assert "[rev]" in txt and "BORRADOR" in txt


def test_pdf_output_or_501(client):
    import shutil
    body = {"instrument": PAGARE, "attorney": ATT, "cutoff_date": "2024-04-20", "format": "pdf"}
    r = client.post("/api/v1/generate-lawsuit", json=body)
    if shutil.which("soffice") or shutil.which("libreoffice"):
        assert r.status_code == 200 and r.content.startswith(b"%PDF")
    else:
        assert r.status_code == 501


def test_pdf_501_when_libreoffice_missing(client, monkeypatch):
    import shutil
    monkeypatch.setattr(shutil, "which", lambda _: None)
    body = {"instrument": PAGARE, "attorney": ATT, "cutoff_date": "2024-04-20", "format": "pdf"}
    assert client.post("/api/v1/generate-lawsuit", json=body).status_code == 501
